"""Test del renderer dell'Excel di avanzamento: confronto con la copia congelata 0.2.0."""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1]
NEW = SCRIPTS / "generate-progress-xlsx.py"
LEGACY = SCRIPTS / "tests/legacy/generate_progress_xlsx_v020.py"


def render(script, plan_dir, out, plan_id=None):
    """Lancia `render` con i file del piano che esistono e ritorna il Workbook letto da `out`."""
    cmd = [sys.executable, str(script), "render"]
    for flag, name in (("--tasks", "TASKS.md"), ("--progress", "PROGRESS.md"), ("--plan", "PLAN.md")):
        if (plan_dir / name).exists():
            cmd += [flag, str(plan_dir / name)]
    if plan_id is not None:
        cmd += ["--plan-id", plan_id]
    cmd += ["--out", str(out), "--today", "2026-10-29"]
    result = subprocess.run(cmd, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    return load_workbook(out)


def sheet_rows(wb, name):
    return list(wb[name].iter_rows(values_only=True))


def legacy_to_float(x):
    """La `to_float` del renderer 0.2.0."""
    return float(re.search(r"[-+]?\d*\.?\d+", str(x).replace(",", ".")).group(0))


def test_clean_fixture_matches_legacy_except_explained(tmp_path, plans_fixtures):
    d = plans_fixtures / "flat/plans/in-progress/2026-09-01_alpha"
    old = render(LEGACY, d, tmp_path / "old.xlsx"); new = render(NEW, d, tmp_path / "new.xlsx")
    o, n = sheet_rows(old, "Task"), sheet_rows(new, "Task")
    eff = o[0].index("Effort")
    for ro, rn in zip(o[1:], n[1:]):
        assert ro[:eff] + ro[eff + 1:] == rn[:eff] + rn[eff + 1:]
        assert legacy_to_float(ro[eff]) == rn[eff]          # «`2 gg`» → 2.0: effort riletto
    assert len(o) == len(n)
    assert sheet_rows(old, "Per Sviluppatore") == sheet_rows(new, "Per Sviluppatore")
    ro, rn = sheet_rows(old, "Riepilogo"), sheet_rows(new, "Riepilogo")
    assert rn[:len(ro)] == ro and any(r[0] == "QUALITÀ DELLA LETTURA" for r in rn[len(ro):])


def test_nested_subplan_rows_recovered(tmp_path, plans_fixtures):
    d = plans_fixtures / "nested/plans/in-progress/2026-07-29_wave-x/sub-a"
    ids = [r[0] for r in sheet_rows(render(NEW, d, tmp_path / "a.xlsx"), "Task")[1:]]
    assert ids == sorted(["T-001", "T-002", "T-003", "T-004", "T-051-FIX", "T-010", "T-011", "T-099"])


def test_unreadable_effort_and_unknown_progress_are_blank(tmp_path, plans_fixtures):
    wb = render(NEW, plans_fixtures / "nested/plans/in-progress/2026-07-29_wave-x/sub-a", tmp_path / "a.xlsx")
    rows = {r[0]: r for r in sheet_rows(wb, "Task")[1:]}; h = sheet_rows(wb, "Task")[0]
    assert rows["T-004"][h.index("Effort")] is None and rows["T-011"][h.index("Effort")] is None
    assert rows["T-004"][h.index("Progresso")] in (None, "")
    summary = dict(sheet_rows(wb, "Riepilogo"))
    assert summary["Effort non leggibile o non dichiarato (task)"] == 3   # T-004, T-011, T-099 (solo nel PROGRESS)
    assert summary["Righe viste / lette"] == "14 / 14"


def test_owner_from_progress_used(tmp_path, plans_fixtures):
    wb = render(NEW, plans_fixtures / "nested/plans/in-progress/2026-07-29_wave-x/sub-a", tmp_path / "a.xlsx")
    h = sheet_rows(wb, "Task")[0]
    assert {r[0]: r for r in sheet_rows(wb, "Task")[1:]}["T-002"][h.index("Owner")] == "Bruno"


def test_placeholder_note_keeps_manual_note(tmp_path, plans_fixtures):
    d = plans_fixtures / "nested/plans/in-progress/2026-07-29_wave-x/sub-a"   # T-003 ha Note «—» nel PROGRESS
    out = tmp_path / "p.xlsx"
    render(NEW, d, out)
    wb = load_workbook(out); ws = wb["Task"]; h = [c.value for c in ws[1]]
    next(r for r in ws.iter_rows(min_row=2) if r[0].value == "T-003")[h.index("Note")].value = "nota del TL"
    wb.save(out)
    wb2 = render(NEW, d, out)
    h2 = sheet_rows(wb2, "Task")[0]
    assert {r[0]: r for r in sheet_rows(wb2, "Task")[1:]}["T-003"][h2.index("Note")] == "nota del TL"


def test_anti_zeroing_guard_still_exits_3(tmp_path, plans_fixtures):
    out = tmp_path / "p.xlsx"
    render(NEW, plans_fixtures / "flat/plans/in-progress/2026-09-01_alpha", out)
    empty = tmp_path / "TASKS.md"; empty.write_text("", encoding="utf-8")
    r = subprocess.run([sys.executable, str(NEW), "render", "--tasks", str(empty), "--out", str(out),
                        "--today", "2026-10-29"], capture_output=True, text=True)
    assert r.returncode == 3


def test_placeholder_owner_dash_kept_name_placeholders_empty(tmp_path):
    """`—` stays `—` as in 0.2.0; `agent` is not an owner: empty cell, grouped under `—`."""
    d = tmp_path / "plan"; d.mkdir()
    (d / "TASKS.md").write_text(
        "| ID | Owner | Effort |\n|---|---|---|\n"
        "| T-001 | — | 1 |\n| T-002 | agent | 2 |\n| T-003 | Alice | 3 |\n",
        encoding="utf-8",
    )
    wb = render(NEW, d, tmp_path / "o.xlsx")
    h = sheet_rows(wb, "Task")[0]
    owners = {r[0]: r[h.index("Owner")] for r in sheet_rows(wb, "Task")[1:]}
    assert owners == {"T-001": "—", "T-002": None, "T-003": "Alice"}
    dev = sheet_rows(wb, "Per Sviluppatore")
    by_owner = {r[0]: r for r in dev[1:]}
    assert set(by_owner) == {"—", "Alice"}
    assert by_owner["—"][dev[0].index("Task totali")] == 2
    assert by_owner["Alice"][dev[0].index("Task totali")] == 1


def test_attenzione_hint_names_the_launcher_and_the_plan_id(tmp_path):
    """Con righe non lette il Riepilogo indica il comando di diagnosi: il launcher, non `-m sdlc_plans`."""
    d = tmp_path / "plan"; d.mkdir()
    (d / "TASKS.md").write_text(
        "| ID | Owner | Effort |\n|---|---|---|\n| T-001 | A | 1 |\n| T-002 | B | 1 |\n"
        "\n### Wave 2\n| T-003 | C | 1 |\n",
        encoding="utf-8",
    )   # T-003 sta fuori dalla tabella: riga vista e non letta
    plan_id = "plans/in-progress/wave-1/sotto-piano"
    wb = render(NEW, d, tmp_path / "d.xlsx", plan_id=plan_id)
    attenzione = dict(sheet_rows(wb, "Riepilogo"))["ATTENZIONE"]
    assert attenzione.startswith("righe non lette: 1")
    assert f"sdlc-plans.py read --plan {plan_id}" in attenzione
    assert "python -m sdlc_plans" not in attenzione
